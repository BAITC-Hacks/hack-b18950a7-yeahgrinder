"""Read partner workbooks by header signatures; never depend on filenames."""
from dataclasses import dataclass, field, fields
from datetime import date
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import tempfile

from dotenv import load_dotenv
import numpy as np
from openpyxl import load_workbook
import pandas as pd

from engine.models import ROOT

load_dotenv(ROOT / ".env")
CACHE_VERSION = 1
MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
TABLES = {
    "sku_master": ["supplier", "sku", "name", "unit", "article", "moq", "moq_source", "flags", "group"],
    "sales_monthly": ["supplier", "sku", "month", "qty_raw", "is_partial"],
    "sales_lines": ["supplier", "sku", "date", "doc_no", "qty"],
    "returns": ["supplier", "sku", "date", "doc_no", "qty"],
    "stock_monthly": ["supplier", "sku", "month", "opening", "closing"],
    "stock_current": ["supplier", "sku", "free_stock", "reserved", "source", "as_of"],
    "transit": ["supplier", "sku", "order_id", "qty", "eta"],
    "seasonality": ["supplier", "month_num", "coef"],
}


class DataError(ValueError):
    pass


@dataclass
class Dataset:
    sku_master: pd.DataFrame
    sales_monthly: pd.DataFrame
    sales_lines: pd.DataFrame
    returns: pd.DataFrame
    stock_monthly: pd.DataFrame
    stock_current: pd.DataFrame
    transit: pd.DataFrame
    seasonality: pd.DataFrame
    as_of: date
    warnings: list[str] = field(default_factory=list)
    fingerprint: str = ""

    def tables(self):
        return {name: getattr(self, name) for name in TABLES}


def text(value):
    return "" if pd.isna(value) else str(value).replace("\xa0", " ").strip()


def number(series):
    return pd.to_numeric(series.astype(str).str.replace(r"[\s\xa0]", "", regex=True)
                         .str.replace(",", ".", regex=False), errors="coerce")


def month_column(value):
    match = re.match(r"([а-яА-Я]+)\.?\s+(\d{4})", text(value))
    if match and match[1].lower()[:3] in MONTHS:
        return pd.Timestamp(int(match[2]), MONTHS[match[1].lower()[:3]], 1)
    return None


def signature(rows):
    for i, row in enumerate(rows[:5]):
        h = [text(v).lower() for v in row]
        if {"дата", "номер", "код", "количество"} <= set(h):
            return "sales_lines", i
        if "свободный остаток" in h and "код 1с" in h:
            return "transit", i
        if any("поступление до" in c for c in h):
            return "transit", i
        if "номенклатура.код" in h and any(month_column(c) is not None for c in h):
            is_stock = any(c in {"ед.", "ед.изм"} for c in h)
            return ("stock_monthly" if is_stock else "sales_monthly"), i
        if "мин. разр. к отгр." in h or ("кратность" in h and "номенклатура.код" in h):
            return "moq", i
    # Dedicated seasonality sheets have their headers further down.
    if any(text(v).lower() in {"сезонность", "норм. коэф."} for row in rows for v in row):
        return "seasonality", 0
    return None


def discover(folder):
    found = {}
    for path in sorted(Path(folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    rows = list(sheet.iter_rows(min_row=1, max_row=45, values_only=True))
                    sig = signature(rows)
                    if sig:
                        kind, header = sig
                        if kind in found and kind != "seasonality":
                            raise DataError(f"Два источника {kind} в {folder}; оставьте одну актуальную выгрузку")
                        found.setdefault(kind, (path, sheet.title, header))
            finally:
                book.close()
        except DataError:
            raise
        except Exception as exc:
            raise DataError(f"Не удалось прочитать {path.name}: {exc}") from exc
    return found


def read_table(source):
    path, sheet, header = source
    df = pd.read_excel(path, sheet_name=sheet, header=header, dtype=object)
    df.columns = [text(c) for c in df.columns]
    return df


def code_column(df):
    return next(c for c in df if c.lower() in {"код", "код 1с", "номенклатура.код"})


def body(df):
    col = code_column(df)
    df = df.copy()
    df["sku"] = df[col].map(text)
    return df[~df.sku.str.lower().isin(["", "итого", "всего", "nan", "none"])].copy()


def item_rows(df, supplier, source):
    name = next((c for c in df if c.lower() in {"номенклатура", "наименование"}), None)
    unit = next((c for c in df if c.lower() in {"ед.", "ед.изм"}), None)
    article = next((c for c in df if "артикул" in c.lower()), None)
    moq = next((c for c in df if c.lower() in {"кратность", "мин. разр. к отгр."}), None)
    out = pd.DataFrame({"sku": df.sku, "supplier": supplier})
    out["name"] = df[name].map(text) if name else ""
    out["unit"] = df[unit].map(text) if unit else ""
    out["article"] = df[article].map(text) if article else ""
    out["moq"] = number(df[moq]) if moq else np.nan
    out["moq_source"] = source if moq else ""
    return out


def read_lines(df, supplier, as_of):
    df = body(df)
    dates = pd.to_datetime(df["Дата"], format="mixed", dayfirst=True, errors="coerce")
    qty = number(df["Количество"])
    mask = dates.notna() & (dates.dt.date <= as_of) & qty.notna()
    if "Документ" in df:
        mask &= df["Документ"].map(text).str.startswith("Расходная накладная")
    if "Склад" in df:
        mask &= df["Склад"].map(text).eq("Алматы")
    lines = pd.DataFrame({"supplier": supplier, "sku": df.sku, "date": dates.dt.normalize(),
                          "doc_no": df["Номер"].map(text), "qty": qty})[mask]
    return lines[lines.qty > 0].reset_index(drop=True), lines[lines.qty < 0].reset_index(drop=True)


def wide(df, supplier, value):
    df = body(df)
    rows = [pd.DataFrame({"supplier": supplier, "sku": df.sku, "month": month_column(c),
                          value: number(df[c]).fillna(0)}) for c in df if month_column(c) is not None]
    return pd.concat(rows, ignore_index=True).groupby(["supplier", "sku", "month"], as_index=False)[value].sum()


def read_season(source, supplier):
    path, sheet, _ = source
    frame = pd.read_excel(path, sheet_name=sheet, header=None)
    candidates = []
    for i, row in frame.iterrows():
        labels = [text(c).lower() for c in row]
        for label in ("сезонность", "норм. коэф."):
            if label in labels:
                col = labels.index(label)
                values = number(frame.iloc[i+1:i+13, col]).to_numpy()
                if len(values) == 12 and np.isfinite(values).all() and (values > 0).all():
                    candidates.append(values)
    if not candidates:
        raise DataError(f"Не найдены 12 положительных коэффициентов сезонности: {path.name}")
    values = candidates[-1]
    return pd.DataFrame({"supplier": supplier, "month_num": range(1, 13), "coef": values / values.mean()})


def read_transit(df, supplier, as_of):
    df = body(df)
    batches = []
    for col in df:
        match = re.search(r"поступление до (\d{2}\.\d{2}\.\d{4})", col, re.I)
        short = re.search(r"в пути\s+(\d{2}\.\d{2})(?:\.(\d{4}))?", col, re.I)
        if not (match or short):
            continue
        eta = pd.to_datetime(match[1] if match else f"{short[1]}.{short[2] or as_of.year}", format="%d.%m.%Y")
        part = pd.DataFrame({"supplier": supplier, "sku": df.sku, "order_id": col,
                             "qty": number(df[col]).fillna(0), "eta": eta})
        batches.append(part[part.qty > 0])
    transit = pd.concat(batches, ignore_index=True) if batches else pd.DataFrame(columns=TABLES["transit"])
    current = pd.DataFrame(columns=TABLES["stock_current"])
    if "Свободный остаток" in df:
        free = number(df["Свободный остаток"])
        if free.isna().any():
            raise DataError("В текущем остатке есть пустые/нечисловые значения")
        current = pd.DataFrame({"supplier": supplier, "sku": df.sku, "free_stock": free.clip(lower=0),
                                "reserved": number(df["Зарезервировано"]).fillna(0),
                                "source": "manager_file", "as_of": pd.Timestamp(as_of)})
        if current.sku.duplicated().any():
            raise DataError("Повтор кода 1С в текущих остатках")
    return transit, current


def supplier_dataset(folder, supplier, as_of):
    sources = discover(folder)
    required = {"sales_lines", "sales_monthly", "stock_monthly", "transit", "seasonality"}
    missing = required - sources.keys()
    if missing:
        raise DataError(f"{supplier}: в {folder} не найдены источники {', '.join(sorted(missing))}")
    frames = {k: read_table(v) for k, v in sources.items() if k != "seasonality"}
    lines, returns = read_lines(frames["sales_lines"], supplier, as_of)
    monthly = wide(frames["sales_monthly"], supplier, "qty_raw")
    monthly = monthly[monthly.month <= pd.Timestamp(as_of)]
    monthly["is_partial"] = monthly.month.eq(pd.Timestamp(as_of).replace(day=1)) & (as_of.day < pd.Timestamp(as_of).days_in_month)
    stock = wide(frames["stock_monthly"], supplier, "opening")
    closing = stock.rename(columns={"opening": "closing"}).copy()
    closing["month"] = closing.month - pd.offsets.MonthBegin(1)
    stock = stock.merge(closing, on=["supplier", "sku", "month"], how="left")
    transit, current = read_transit(frames["transit"], supplier, as_of)
    # Priority: dedicated MOQ and transit descriptions (coil instructions), then other sources.
    parts = [item_rows(body(frames[k]), supplier, k) for k in
             ["moq", "transit", "sales_monthly", "stock_monthly", "sales_lines"] if k in frames]
    candidates = pd.concat(parts, ignore_index=True).replace({"": None})
    candidates.loc[candidates.moq <= 0, ["moq", "moq_source"]] = [np.nan, None]
    master = candidates.groupby(["supplier", "sku"], as_index=False, sort=True).first()
    master["flags"] = master.moq.map(lambda v: ["MOQ_MISSING"] if pd.isna(v) else [])
    master["moq"] = np.ceil(master.moq.fillna(1)).astype(int)
    master["moq_source"] = master.moq_source.fillna("default")
    master["name"] = master.name.fillna(master.sku)
    master["unit"] = master.unit.fillna("шт")
    master["group"] = master.sku.str[:4]
    coil = master.name.str.contains(r"305\s*м", case=False, regex=True) & master.name.str.contains("БУХТАМИ", case=False)
    for i in master.index[coil]:
        master.at[i, "moq"] = 305
        master.at[i, "moq_source"] = "coil_instruction"
        master.at[i, "flags"] = master.at[i, "flags"] + ["COIL_305"]
    warnings = []
    # No current snapshot for a SKU: opening of this month minus sales since, flagged as estimate.
    unknown = set(master.sku) - set(current.sku)
    if unknown:
        opening = stock[stock.month.eq(pd.Timestamp(as_of).replace(day=1))].set_index("sku").opening
        sold = lines[lines.date >= pd.Timestamp(as_of).replace(day=1)].groupby("sku").qty.sum()
        estimated = master.loc[master.sku.isin(unknown), ["supplier", "sku"]].copy()
        estimated["free_stock"] = (estimated.sku.map(opening).fillna(0) - estimated.sku.map(sold).fillna(0)).clip(lower=0)
        estimated["reserved"] = 0.0
        estimated["source"] = "estimated_opening_minus_sales"
        estimated["as_of"] = pd.Timestamp(as_of)
        current = pd.concat([current, estimated], ignore_index=True) if len(current) else estimated
        master["flags"] = [f + (["STOCK_ESTIMATED"] if sku in unknown else []) for sku, f in zip(master.sku, master["flags"])]
        warnings.append(f"{supplier}: для {len(unknown)} SKU текущий остаток оценён как остаток на 1-е число "
                        "минус продажи месяца (поступления не видны); подтвердить в 1С.")
    absent = set(master.sku) - set(monthly.sku)
    if absent:
        fallback = lines[lines.sku.isin(absent)].copy()
        fallback["month"] = fallback.date.dt.to_period("M").dt.to_timestamp()
        fallback = fallback.groupby(["supplier", "sku", "month"], as_index=False).qty.sum().rename(columns={"qty": "qty_raw"})
        fallback["is_partial"] = fallback.month.eq(pd.Timestamp(as_of).replace(day=1)) & (as_of.day < pd.Timestamp(as_of).days_in_month)
        monthly = pd.concat([monthly, fallback], ignore_index=True)
        master["flags"] = [f + (["MONTHLY_MISSING"] if sku in absent else []) for sku, f in zip(master.sku, master["flags"])]
        warnings.append(f"{supplier}: {len(absent)} SKU отсутствуют в месячном отчёте; использованы доступные накладные.")
    return Dataset(master[TABLES["sku_master"]], monthly, lines, returns, stock, current, transit,
                   read_season(sources["seasonality"], supplier), as_of, warnings)


def data_paths():
    root = Path(os.environ.get("DATA_DIR", ROOT / "data/raw"))
    return {"IEK": Path(os.environ.get("IEK_DATA_DIR", root / "IEK")),
            "Systeme Electric": Path(os.environ.get("SE_DATA_DIR", root / "SE"))}


def load(paths=None, as_of=date(2026, 9, 22), use_cache=True, cache_dir=None):
    as_of = pd.Timestamp(as_of).date()
    paths = paths or data_paths()
    paths = {k: Path(v).expanduser().resolve() for k, v in paths.items()}
    if set(paths) - {"IEK", "Systeme Electric"}:
        raise DataError("Неизвестный поставщик")
    for supplier, folder in paths.items():
        if not folder.is_dir():
            raise DataError(f"{supplier}: нет папки {folder}. Укажите IEK_DATA_DIR / SE_DATA_DIR в .env")
    files = [(s, str(p), p.stat().st_size, p.stat().st_mtime_ns) for s, folder in sorted(paths.items())
             for p in sorted(folder.glob("*.xlsx")) if not p.name.startswith("~$")]
    fingerprint = sha256(json.dumps([CACHE_VERSION, str(as_of), files]).encode()).hexdigest()
    cache_root = Path(cache_dir or ROOT / "data/processed")
    cache = cache_root / fingerprint
    if use_cache and (cache / "meta.json").exists():
        try:
            meta = json.loads((cache / "meta.json").read_text())
            tables = {k: pd.read_parquet(cache / f"{k}.parquet") for k in TABLES}
            tables["sku_master"]["flags"] = tables["sku_master"]["flags"].map(list)
            return Dataset(**tables, as_of=as_of, warnings=meta["warnings"], fingerprint=fingerprint)
        except (OSError, ValueError):
            pass
    datasets = [supplier_dataset(folder, supplier, as_of) for supplier, folder in paths.items()]
    tables = {k: pd.concat([getattr(ds, k) for ds in datasets], ignore_index=True) for k in TABLES}
    ds = Dataset(**tables, as_of=as_of, warnings=[w for d in datasets for w in d.warnings], fingerprint=fingerprint)
    if use_cache:
        cache_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=cache_root) as temp:
            staging = Path(temp)
            for k, df in ds.tables().items():
                df.to_parquet(staging / f"{k}.parquet", index=False)
            (staging / "meta.json").write_text(json.dumps({"warnings": ds.warnings}, ensure_ascii=False))
            cache.mkdir(exist_ok=True)
            # Manifest last: concurrent readers only consume complete snapshots.
            for p in sorted(staging.iterdir(), key=lambda p: p.name == "meta.json"):
                os.replace(p, cache / p.name)
    return ds
