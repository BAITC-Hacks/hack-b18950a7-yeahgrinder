"""Integration regressions for the supplied September 2026 partner archives.

No partner data is checked into git. See docs/TESTING.md for required CLI arguments.
"""
from io import BytesIO
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import numpy as np
import pandas as pd
import pytest

import api
import web_data
from engine.loader import discover
from engine.models import Params
from engine.run import compute
from engine.validate import validate_orders
from service import ProcurementService

pytestmark = pytest.mark.real_data


def test_all_six_sources_discovered_for_each_supplier(partner_inputs):
    root, extracted = partner_inputs
    required = {'sales_lines', 'sales_monthly', 'stock_monthly', 'transit', 'moq', 'seasonality'}
    assert {supplier: len(files) for supplier, files in extracted.items()} == {'IEK': 6, 'Systeme Electric': 6}
    for supplier in ('IEK', 'SE'):
        assert set(discover(root / supplier)) == required, f'Потерян источник {supplier}'


def test_loaded_snapshot_is_complete(partner_dataset):
    ds = partner_dataset
    assert ds.as_of == pd.Timestamp('2026-09-22')
    assert set(ds.items.supplier) == {'ИЭК', 'Systeme Electric'}
    assert ds.items.groupby('supplier').size().to_dict() == {'ИЭК': 2860, 'Systeme Electric': 701}
    assert not ds.items.code.duplicated().any()
    assert ds.items.code.map(lambda c: isinstance(c, str) and c == c.strip()).all()
    assert ds.lines.qty.gt(0).all(), 'Возвраты не должны уменьшать ряд положительного спроса'
    assert set(ds.stock.month.dt.to_period('M').astype(str)) == set(pd.period_range('2024-01', '2026-09', freq='M').astype(str))
    missing = [w for w in ds.warnings if any(x in w for x in ('нет файла', 'не удалось прочитать', 'пропущен:'))]
    assert not missing, '\n'.join(missing)
    assert ds.items.moq.gt(0).all()
    for supplier, season in ds.supplier_season.items():
        assert len(season) == 12 and np.isfinite(season).all() and season.gt(0).all(), supplier
        assert float(season.mean()) == pytest.approx(1, abs=.02)
    cable = ds.items.loc[ds.items.code.eq('200400085_')].iloc[0]
    assert cable.unit == 'м'


def test_loop_raw_invoice_is_preserved(partner_dataset):
    rows = partner_dataset.lines.query("code == '130200305_' and qty == 210000")
    assert len(rows) == 1
    row = rows.iloc[0]
    assert row.date.normalize() == pd.Timestamp('2025-06-09')
    assert '20000064179' in row.doc


def test_every_loaded_item_has_valid_explained_order(partner_dataset, partner_result):
    orders = partner_result.orders
    assert set(orders.sku) == set(partner_dataset.items.code), 'Расчёт не должен молча терять товары'
    assert len(orders) == len(partner_dataset.items)
    assert not validate_orders(orders)
    assert orders.reason_text.str.strip().str.len().gt(0).all()
    assert set(orders.urgency) <= {'CRITICAL', 'HIGH', 'PLANNED', 'OK'}
    assert (orders.recommended_qty >= 0).all()
    assert (orders.recommended_qty % orders.moq == 0).all()
    for summary in partner_result.summary:
        part = orders.loc[orders.supplier.eq(summary['supplier'])]
        assert summary['sku_count'] == len(part)
        assert summary['critical_count'] == part.urgency.eq('CRITICAL').sum()
        assert summary['order_count'] == part.recommended_qty.gt(0).sum()


def test_real_anomalies_and_stockouts_are_explained(partner_result):
    loop = partner_result.orders.loc[partner_result.orders.sku.eq('130200305_')].iloc[0]
    assert 'ONE_OFF_INVOICE' in loop.reason_codes
    event = next(e for e in loop.excluded_events if e.get('raw_qty') == 210000)
    assert event['month'] == '2025-06'
    assert event['excluded_qty'] == 210000  # current main excludes the whole one-off line
    assert any('STOCKOUT_RESTORED' in c for c in partner_result.orders.reason_codes)
    restored = [e for events in partner_result.orders.restored_events for e in events]
    # Historical returns may be negative; restoration can bring the month to zero.
    assert restored and all(e['qty'] > 0 and e['clean_qty'] >= 0 for e in restored)
    # One month may contain BOTH an excluded spike and restored regular demand.
    # Compare to demand after exclusion, not to the uncorrected (possibly huge) sale.
    history = partner_result.history.loc[partner_result.history.restored_qty.gt(0)]
    before_restore = history.qty - history.excluded_qty
    assert (history.clean_qty > before_restore).all()
    np.testing.assert_allclose(history.clean_qty, before_restore + history.restored_qty, atol=1e-6)


def test_growth_scenario_and_ui_payload_on_real_sources(partner_dataset, partner_result, monkeypatch):
    grown = compute(partner_dataset, Params(growth_pct=20))
    baseline = partner_result.orders.set_index('sku').sort_index()
    growth = grown.orders.set_index('sku').sort_index()
    # Both public values are rounded to 0.1: combined absolute error <= .05 + 1.2*.05.
    np.testing.assert_allclose(growth.forecast_horizon, baseline.forecast_horizon * 1.2, rtol=0, atol=.110001)
    assert (growth.recommended_qty >= baseline.recommended_qty).all()
    assert (growth.recommended_qty > baseline.recommended_qty).any()
    # The endpoint still serializes the actual adapter output, with calculations shared
    # only to avoid running the same expensive full batch three times in one test.
    monkeypatch.setattr(web_data, 'load', lambda: partner_dataset)
    monkeypatch.setattr(web_data, 'compute', lambda ds, p: grown if p.growth_pct else partner_result)
    web_data.payload.cache_clear()
    try:
        response = TestClient(api.app).get('/ui/data')
        assert response.status_code == 200, response.text[:1000]
        payload = response.json()
        json.dumps(payload, allow_nan=False)
        assert payload['meta']['live'] is True
        assert len(payload['products']) == len(baseline)
        assert len({p['id'] for p in payload['products']}) == len(baseline)
        for p in payload['products']:
            assert p['recommended_qty'] == baseline.at[p['sku'], 'recommended_qty']
            assert p['growthScenario']['recommended_qty'] == growth.at[p['sku'], 'recommended_qty']
            assert len(p['history']) == 8
        loop = next(p for p in payload['products'] if p['sku'] == '130200305_')
        assert 'Архив' in loop['history_label']
        assert any(e['qty'] == 210000 for e in loop['excluded_events'])
    finally:
        web_data.payload.cache_clear()


def test_api_approval_export_and_what_if_on_partner_data(partner_dataset, tmp_path, monkeypatch):
    service = ProcurementService(storage=tmp_path, dataset=partner_dataset)
    monkeypatch.setattr(api, 'service', lambda: service)
    client = TestClient(api.app)
    run = client.post('/runs')
    assert run.status_code == 200, run.text[:1000]
    run_id = run.json()['run_id']
    page = client.get(f'/runs/{run_id}/orders', params={'supplier': 'IEK', 'only_orders': True})
    assert page.status_code == 200
    selected = page.json()['rows'][0]
    sku = selected['sku']
    edited_qty = selected['recommended_qty'] + selected['moq']
    approval = client.post(f'/runs/{run_id}/decisions', json={'supplier': 'IEK', 'decision': 'approve',
                           'edits': {sku: edited_qty}, 'reason': 'Integration test'})
    assert approval.status_code == 200, approval.text[:1000]
    result = client.get(f"/approvals/{approval.json()['approval_id']}/export")
    assert result.status_code == 200
    with BytesIO(result.content) as stream:
        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            exported = list(book['Заказ'].iter_rows(min_row=2, values_only=True))
            assert exported and all(r[4] > 0 and r[4] % r[5] == 0 for r in exported)
            assert next(r[4] for r in exported if r[0] == sku) == edited_qty
            assert 'Параметры и допущения' in book.sheetnames
        finally:
            book.close()
    # A hypothetical arrival reduces need without changing the saved run.
    delta = client.post(f'/runs/{run_id}/what-if', json={'supplier': 'IEK', 'sku': sku, 'extra_transit': int(selected['moq']) * 10})
    assert delta.status_code == 200, delta.text[:1000]
    assert delta.json()['saved'] is False
    assert delta.json()['changes'][0]['delta'] <= 0
    unchanged = client.get(f'/runs/{run_id}/orders/IEK/{sku}').json()['order']
    assert unchanged['recommended_qty'] == selected['recommended_qty']
    assert client.post(f'/runs/{run_id}/chat', json={'message': 'Почему столько?'}).status_code == 503
