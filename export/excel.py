from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font


def safe_text(value):
    # Spreadsheet formula injection protection, including names from source Excel.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def export_xlsx(rows, params, warnings, approval_id):
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказ"
    ws.append(["Код 1с", "Артикул", "Наименование", "Ед.", "Кол-во", "Кратность", "Срочность", "Обоснование"])
    for row in rows:
        if row['recommended_qty'] > 0:
            ws.append([safe_text(row[k]) for k in ['sku', 'article', 'name', 'unit', 'recommended_qty', 'moq', 'urgency', 'reason_text']])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in {'A':18, 'B':24, 'C':65, 'D':10, 'E':15, 'F':15, 'G':18, 'H':100}.items():
        ws.column_dimensions[col].width = width
    notes = wb.create_sheet("Параметры и допущения")
    notes.append(["Утверждение", approval_id])
    for key, value in params.items():
        notes.append([key, safe_text(json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value))])
    for warning in warnings:
        notes.append(["Допущение", safe_text(warning)])
    notes.column_dimensions['A'].width = 25
    notes.column_dimensions['B'].width = 110
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
